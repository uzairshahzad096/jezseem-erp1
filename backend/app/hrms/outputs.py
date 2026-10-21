"""HR downloads: Excel sheets and the payslip PDF. Built from rows the
screens already show, so a file never disagrees with the screen."""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.documents.pdf import INK, LINE, _text
from app.reports.stock_workbook import _table, _title

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY = "#,##0.00"
QTY = "0.##"
DATE = "DD-MMM-YYYY"
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
          "October", "November", "December"]


def sheet(title: str, subtitle: str, columns: list[tuple], rows: list[dict],
          totals: dict | None = None) -> bytes:
    """One styled sheet. columns: (key, label, width, number format or None)."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    stamp = f"{subtitle} · generated {datetime.now():%d-%b-%Y %H:%M}"
    header = _title(ws, title, stamp, len(columns))
    _table(ws, header, columns, rows)
    if totals and rows:
        row = header + len(rows) + 1
        ws.cell(row=row, column=1, value="Total").font = Font(bold=True)
        for c, (key, _label, _width, fmt) in enumerate(columns, start=1):
            if key in totals:
                cell = ws.cell(row=row, column=c, value=totals[key])
                cell.font = Font(bold=True)
                if fmt:
                    cell.number_format = fmt
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def _payslip_story(company: str, run: dict, slip: dict) -> list:
    base = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=14, textColor=INK)
    label = ParagraphStyle("l", parent=base["Normal"], fontName="Helvetica", fontSize=8, textColor=colors.grey)
    value = ParagraphStyle("v", parent=base["Normal"], fontName="Helvetica", fontSize=9.5, textColor=INK)
    bold = ParagraphStyle("b", parent=value, fontName="Helvetica-Bold")
    small = ParagraphStyle("s", parent=value, fontSize=8, textColor=colors.grey)

    def cell(caption, text, style=value):
        return [Paragraph(caption, label), Paragraph(_text(str(text) if text not in (None, "") else "—"), style)]

    period = f"{MONTHS[run['period_month'] - 1]} {run['period_year']}"
    story = [Paragraph(_text(company), bold), Spacer(1, 3),
             Paragraph(f"PAYSLIP — {period}", title), Spacer(1, 9)]
    facts = [
        cell("Employee", f"{slip['employee_code']}  {slip['employee_name']}", bold) + cell("Payroll", run["run_number"]),
        cell("Department", slip["department"]) + cell("Designation", slip["designation"]),
        cell("Payment", slip["payment_mode"].title()) + cell("Bank account", " ".join(
            x for x in (slip["bank_name"], slip["bank_account_number"]) if x)),
        cell("Days in month", slip["days_in_period"]) + cell("Present / paid leave",
                                                              f"{slip['present_days']:g} / {slip['paid_leave_days']:g}"),
        cell("Absent / unpaid leave", f"{slip['absent_days']:g} / {slip['unpaid_leave_days']:g}")
        + cell("Late arrivals / overtime hours", f"{slip['late_count']} / {slip['overtime_hours']:g}"),
    ]
    table = Table(facts, colWidths=[34 * mm, 53 * mm, 40 * mm, 47 * mm])
    table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story += [table, Spacer(1, 8)]

    earnings = [x for x in slip["lines"] if x["kind"] == "EARNING"]
    deductions = [x for x in slip["lines"] if x["kind"] == "DEDUCTION"]
    rows = [["Earnings", "Amount", "Deductions", "Amount"]]
    for i in range(max(len(earnings), len(deductions), 1)):
        e = earnings[i] if i < len(earnings) else None
        d = deductions[i] if i < len(deductions) else None
        rows.append([Paragraph(_text(e["name"]), value) if e else "", f"{e['amount']:,.2f}" if e else "",
                     Paragraph(_text(d["name"]), value) if d else "", f"{d['amount']:,.2f}" if d else ""])
    rows.append(["Total earnings", f"{slip['total_earnings']:,.2f}",
                 "Total deductions", f"{slip['total_deductions']:,.2f}"])
    lines = Table(rows, colWidths=[57 * mm, 30 * mm, 57 * mm, 30 * mm])
    lines.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, LINE), ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFF4EF")),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("ALIGN", (3, 0), (3, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    net = Table([["NET PAY", f"PKR {slip['net_pay']:,.2f}"]], colWidths=[117 * mm, 57 * mm])
    net.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 11),
        ("ALIGN", (1, 0), (1, 0), "RIGHT"), ("BOX", (0, 0), (-1, -1), 0.8, INK),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story += [lines, Spacer(1, 8), net, Spacer(1, 6)]
    if run["status"] not in ("APPROVED", "PAID"):
        story.append(Paragraph(f"{run['status']} — not an approved payslip.", small))
    elif run.get("payment_date"):
        story.append(Paragraph(f"Paid on {run['payment_date']:%d-%b-%Y}.", small))
    story += [Spacer(1, 26), Table([["Prepared by", "Approved by", "Received by (employee)"]],
                                   colWidths=[58 * mm] * 3,
                                   style=TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.5, INK),
                                                     ("FONTSIZE", (0, 0), (-1, -1), 8.5)]))]
    return story


def payslips_pdf(company: str, run: dict, slips: list[dict]) -> bytes:
    """One payslip per page."""
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm, title=f"Payslips {run['run_number']}")
    story: list = []
    for index, slip in enumerate(slips):
        if index:
            story.append(PageBreak())
        story += _payslip_story(company, run, slip)
    doc.build(story)
    return stream.getvalue()
