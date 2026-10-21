"""
Opening balances: the accountant's approved position at the cut-over date.

    Import -> Validate -> Preview -> Approve -> Post -> Reconcile -> Complete

Nothing here derives a figure from existing ERP data. Every amount comes
from schedules the accountant prepares from the old system: the trial
balance, the customer and supplier schedules, the inventory valuation and
the bank balances. The sales orders, zero-cost opening stock and demo
procurement records in the database are never read.

How the schedules combine
-------------------------
The GL schedule (the legacy trial balance) is the master and must balance.
Control accounts are then broken down by their subledger schedule:

  RECEIVABLE account  = sum of the AR schedule (one line per customer invoice)
  PAYABLE account     = sum of the AP schedule (one line per supplier bill)
  INVENTORY account   = sum of the inventory valuation (qty × unit cost)
  BANK account        = statement balance + uncleared deposits − uncleared payments

Validation refuses the batch unless each of these ties exactly. The opening
journal then carries the schedule detail in place of the single trial-balance
line: the receivable account holds a line per customer, the inventory account
a line per item. Subledger reports can therefore be built from the ledger
itself, and the journal still balances because the trial balance did.
"""

import csv
import hashlib
import io
import re
import uuid
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from app.audit.service import record, snapshot
from app.companies.models import Branch, FinancialYear
from app.finance import ledger, numbering, periods
from app.finance.context import FinanceContext
from app.finance.errors import FinanceError, db_guard
from app.finance.models import (
    OB_APPROVED, OB_CANCELLED, OB_COMPLETED, OB_DRAFT, OB_POSTED, OB_VALIDATED, OPENING,
    RECONCILIATION_CATEGORIES, SCHEDULES, Account, JournalEntry, JournalEntryLine,
    OpeningBalanceBatch, OpeningBalanceImport, OpeningBalanceLine, OpeningBalanceReconciliation,
)
from app.finance.posting import PostingLine, PostingRequest, money, post as post_event
from app.inventory.models import InventoryItem, ItemLocation
from app.partners.models import BusinessPartner

ZERO = Decimal("0")

#: Which control type each subledger schedule breaks down.
SCHEDULE_CONTROL = {"AR": "RECEIVABLE", "AP": "PAYABLE", "INVENTORY": "INVENTORY", "BANK": "BANK"}
#: Further control types a schedule may carry. Customer advances (Phase 2)
#: are entered per customer in the Customers schedule, as negative amounts
#: against the advances account, because that account needs a customer on
#: every line just like the receivable account.
EXTRA_CONTROL = {"CUSTOMER_ADVANCE": "AR"}
ALL_CONTROLS = set(SCHEDULE_CONTROL.values()) | set(EXTRA_CONTROL)


def _schedule_of(control_type: str | None) -> str | None:
    if control_type in EXTRA_CONTROL:
        return EXTRA_CONTROL[control_type]
    return next((s for s, c in SCHEDULE_CONTROL.items() if c == control_type), None)

#: Template columns per schedule, as the accountant sees them.
TEMPLATE_COLUMNS = {
    "GL": ["Account Code", "Account Name", "Debit", "Credit", "Branch", "Reference"],
    "AR": ["Customer Code", "Customer Name", "Invoice Reference", "Invoice Date", "Due Date",
           "Outstanding Amount", "Account Code", "Branch"],
    "AP": ["Supplier Code", "Supplier Name", "Bill Reference", "Bill Date", "Due Date",
           "Outstanding Amount", "Account Code", "Branch"],
    "INVENTORY": ["Item Code", "Warehouse", "Quantity", "Unit Cost", "Total Value",
                  "Account Code", "Branch"],
    "BANK": ["Bank", "Account Number", "Account Code", "Statement Balance",
             "Uncleared Deposits", "Uncleared Payments", "Branch"],
}
#: Sheet names in the template workbook.
SHEETS = {"GL": "Trial Balance", "AR": "Customers", "AP": "Suppliers",
          "INVENTORY": "Inventory", "BANK": "Banks"}

#: Header aliases: normalised header -> field. Normalised = lower case,
#: letters and digits only, so "Invoice Ref." and "invoice_ref" both match.
ALIASES = {
    "GL": {
        "accountcode": "account_code", "code": "account_code", "glcode": "account_code",
        "account": "account_code", "accountname": "account_name", "name": "account_name",
        "description": "account_name", "debit": "debit", "dr": "debit", "credit": "credit",
        "cr": "credit", "branch": "branch_code", "branchcode": "branch_code",
        "reference": "reference", "ref": "reference",
    },
    "AR": {
        "customercode": "party_code", "code": "party_code", "partycode": "party_code",
        "customername": "party_name", "customer": "party_name", "name": "party_name",
        "invoicereference": "reference", "invoiceref": "reference", "invoice": "reference",
        "invoiceno": "reference", "invoicenumber": "reference", "reference": "reference",
        "invoicedate": "document_date", "date": "document_date", "duedate": "due_date",
        "outstandingamount": "amount", "outstanding": "amount", "amount": "amount",
        "balance": "amount", "accountcode": "account_code", "branch": "branch_code",
    },
    "AP": {
        "suppliercode": "party_code", "code": "party_code", "partycode": "party_code",
        "suppliername": "party_name", "supplier": "party_name", "name": "party_name",
        "billreference": "reference", "billref": "reference", "bill": "reference",
        "billno": "reference", "invoicereference": "reference", "reference": "reference",
        "billdate": "document_date", "invoicedate": "document_date", "date": "document_date",
        "duedate": "due_date", "outstandingamount": "amount", "outstanding": "amount",
        "amount": "amount", "balance": "amount", "accountcode": "account_code",
        "branch": "branch_code",
    },
    "INVENTORY": {
        "itemcode": "item_code", "code": "item_code", "item": "item_code",
        "warehouse": "location_name", "location": "location_name", "quantity": "quantity",
        "qty": "quantity", "unitcost": "unit_cost", "cost": "unit_cost", "rate": "unit_cost",
        "totalvalue": "amount", "value": "amount", "amount": "amount",
        "accountcode": "account_code", "branch": "branch_code",
    },
    "BANK": {
        "bank": "bank_name", "bankname": "bank_name", "accountnumber": "bank_account_number",
        "accountno": "bank_account_number", "account": "bank_account_number",
        "accountcode": "account_code", "glaccount": "account_code", "glcode": "account_code",
        "statementbalance": "statement_balance", "uncleareddeposits": "uncleared_deposits",
        "unclearedreceipts": "uncleared_deposits", "depositsintransit": "uncleared_deposits",
        "unclearedpayments": "uncleared_payments", "unpresentedcheques": "uncleared_payments",
        "branch": "branch_code",
    },
}
MONEY_FIELDS = {"debit", "credit", "amount", "statement_balance", "uncleared_deposits",
                "uncleared_payments"}
QTY_FIELDS = {"quantity", "unit_cost"}
DATE_FIELDS = {"document_date", "due_date"}


# ------------------------------------------------------------ parsing

def _norm(header) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def parse_decimal(value, places: int = 2) -> Decimal | None:
    """Accountants' numbers: '1,234.50', '(500)', ' - ', 1234.5."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("not a number")
    if isinstance(value, (int, Decimal)):
        number = Decimal(value)
    elif isinstance(value, float):
        number = Decimal(repr(value))
    else:
        text = str(value).strip().replace(",", "").replace("PKR", "").replace("Rs", "").strip()
        if text in ("", "-", "–"):
            return None
        negative = text.startswith("(") and text.endswith(")")
        text = text.strip("()")
        try:
            number = Decimal(text)
        except InvalidOperation as error:
            raise ValueError(f"'{value}' is not a number") from error
        if negative:
            number = -number
    return number.quantize(Decimal(1).scaleb(-places))


def parse_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{value}' is not a date (use DD-MM-YYYY)")


def _rows_from_table(table: list[list]) -> list[dict]:
    """First non-empty row is the header; blank rows are skipped."""
    rows = [r for r in table if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    return [{str(h).strip(): v for h, v in zip(header, r) if h not in (None, "")} for r in body]


def read_file(file_name: str, content: bytes) -> dict[str, list[dict]]:
    """Rows per sheet name for a workbook, or {"": rows} for a CSV."""
    name = (file_name or "").lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        return {"": _rows_from_table(list(csv.reader(io.StringIO(text))))}
    if name.endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        return {sheet.title: _rows_from_table([list(r) for r in sheet.iter_rows(values_only=True)])
                for sheet in workbook.worksheets}
    raise FinanceError("Upload an Excel workbook (.xlsx) or a CSV file.")


def _sheet_schedule(title: str) -> str | None:
    key = _norm(title)
    for schedule, sheet in SHEETS.items():
        if key in (_norm(sheet), schedule.lower()):
            return schedule
    return {"tb": "GL", "trialbalance": "GL", "gl": "GL", "ar": "AR", "receivables": "AR",
            "ap": "AP", "payables": "AP", "stock": "INVENTORY", "bank": "BANK"}.get(key)


def map_row(schedule: str, raw: dict) -> tuple[dict, list[str]]:
    """File row -> line fields, plus any parse errors."""
    aliases = ALIASES[schedule]
    fields, errors = {}, []
    for header, value in raw.items():
        field = aliases.get(_norm(header))
        if field is None or field in fields:
            continue
        try:
            if field in MONEY_FIELDS:
                fields[field] = parse_decimal(value, 2)
            elif field in QTY_FIELDS:
                fields[field] = parse_decimal(value, 4)
            elif field in DATE_FIELDS:
                fields[field] = parse_date(value)
            else:
                text = str(value).strip() if value is not None else ""
                if isinstance(value, float) and value.is_integer():
                    text = str(int(value))  # Excel turns code 1110 into 1110.0
                fields[field] = text or None
        except ValueError as error:
            errors.append(f"{header}: {error}")
    return fields, errors


def _json_safe(raw: dict) -> dict:
    out = {}
    for key, value in raw.items():
        if isinstance(value, (datetime, date)):
            value = value.isoformat()
        elif isinstance(value, Decimal):
            value = str(value)
        elif value is not None and not isinstance(value, (str, int, float, bool)):
            value = str(value)
        out[str(key)] = value
    return out


# ------------------------------------------------------------ batches

def get_batch(db: Session, ctx: FinanceContext, batch_id: uuid.UUID) -> OpeningBalanceBatch:
    batch = db.scalar(select(OpeningBalanceBatch).options(selectinload(OpeningBalanceBatch.lines))
                      .where(OpeningBalanceBatch.id == batch_id))
    if batch is None or batch.company_id != ctx.company_id:
        raise FinanceError("That opening balance batch does not exist.", 404)
    if ctx.branch_ids is not None and batch.branch_id not in ctx.branch_ids:
        raise FinanceError("That opening balance batch does not exist.", 404)
    return batch


def create_batch(db: Session, ctx: FinanceContext, data: dict) -> OpeningBalanceBatch:
    fy = db.get(FinancialYear, data["financial_year_id"])
    if fy is None or fy.company_id != ctx.company_id:
        raise FinanceError("That financial year does not exist.")
    start, end = periods.fy_dates(fy)
    cutover: date = data["cutover_date"]
    if not start <= cutover <= end:
        raise FinanceError(f"The cut-over date must fall within {fy.name} "
                           f"({start:%d-%b-%Y} to {end:%d-%b-%Y}).")
    branch_id = ctx.default_branch(data.get("branch_id"))
    batch = OpeningBalanceBatch(
        company_id=ctx.company_id, branch_id=branch_id, financial_year_id=fy.id,
        batch_number=numbering.allocate(db, ctx.company_id, "OB", fy),
        cutover_date=cutover, status=OB_DRAFT,
        source_document=data.get("source_document"), notes=data.get("notes"),
        prepared_by=ctx.user_id, created_by=ctx.user_id,
    )
    db.add(batch)
    with db_guard(db):
        db.flush()
    record(db, ctx.audit, action="create", module="finance", entity="opening_balance_batches",
           entity_id=batch.id, company_id=ctx.company_id, new_values=snapshot(batch))
    return batch


def _editable(batch: OpeningBalanceBatch) -> None:
    if batch.status not in (OB_DRAFT, OB_VALIDATED):
        raise FinanceError(f"Batch {batch.batch_number} is {batch.status}; its lines are frozen.")


def import_rows(db: Session, ctx: FinanceContext, batch: OpeningBalanceBatch, schedule: str,
                rows: list[dict], *, file_name: str | None, file_hash: str) -> int:
    """Replace one schedule's lines with these rows. Resolution runs at once
    so the screen can show problems row by row; the batch returns to DRAFT
    until it is validated again."""
    if schedule not in SCHEDULES:
        raise FinanceError(f"Unknown schedule {schedule}.")
    _editable(batch)

    duplicate = db.scalar(
        select(OpeningBalanceBatch.batch_number)
        .join(OpeningBalanceImport, OpeningBalanceImport.batch_id == OpeningBalanceBatch.id)
        .where(OpeningBalanceImport.file_hash == file_hash,
               OpeningBalanceImport.schedule == schedule,
               OpeningBalanceBatch.company_id == ctx.company_id,
               OpeningBalanceBatch.id != batch.id,
               OpeningBalanceBatch.status != OB_CANCELLED))
    if duplicate:
        raise FinanceError(f"This file has already been imported into batch {duplicate}. "
                           "Cancel that batch first if it should be replaced.")

    for line in [l for l in batch.lines if l.schedule == schedule]:
        db.delete(line)
    with db_guard(db):
        db.flush()

    count = 0
    for n, raw in enumerate(rows, start=1):
        fields, errors = map_row(schedule, raw)
        line = OpeningBalanceLine(batch_id=batch.id, schedule=schedule, row_number=n,
                                  raw=_json_safe(raw), errors=errors)
        for key, value in fields.items():
            if key in ("debit", "credit"):
                value = value if value is not None else ZERO
            setattr(line, key, value)
        db.add(line)
        count += 1

    db.add(OpeningBalanceImport(batch_id=batch.id, schedule=schedule, file_name=file_name,
                                file_hash=file_hash, row_count=count, imported_by=ctx.user_id))
    batch.status = OB_DRAFT
    batch.validated_at = None
    batch.updated_by = ctx.user_id
    with db_guard(db):
        db.flush()
    db.refresh(batch)
    resolve(db, ctx, batch)
    record(db, ctx.audit, action="import", module="finance", entity="opening_balance_batches",
           entity_id=batch.id, company_id=ctx.company_id,
           new_values={"schedule": schedule, "rows": count, "file_name": file_name,
                       "file_hash": file_hash})
    return count


def import_file(db: Session, ctx: FinanceContext, batch: OpeningBalanceBatch, schedule: str | None,
                file_name: str, content: bytes) -> dict[str, int]:
    """One CSV for a named schedule, or a workbook whose sheets are schedules
    (the template's sheet names)."""
    if not content:
        raise FinanceError("The file is empty.")
    file_hash = hashlib.sha256(content).hexdigest()
    sheets = read_file(file_name, content)
    imported: dict[str, int] = {}
    if "" in sheets:  # CSV
        if not schedule:
            raise FinanceError("Say which schedule this CSV holds.")
        imported[schedule] = import_rows(db, ctx, batch, schedule, sheets[""],
                                         file_name=file_name, file_hash=file_hash)
        return imported
    for title, rows in sheets.items():
        target = _sheet_schedule(title)
        if target is None or (schedule and target != schedule) or not rows:
            continue
        imported[target] = import_rows(db, ctx, batch, target, rows,
                                       file_name=file_name, file_hash=file_hash)
    if not imported:
        raise FinanceError("No schedule found in the workbook. Use the template's sheet names: "
                           + ", ".join(SHEETS.values()) + ".")
    return imported


# ------------------------------------------------------------ resolution & validation

class _Lookups:
    def __init__(self, db: Session, company_id: uuid.UUID):
        self.accounts = {a.code: a for a in db.scalars(
            select(Account).where(Account.company_id == company_id))}
        partners = db.scalars(select(BusinessPartner).where(BusinessPartner.company_id == company_id)).all()
        self.partner_by_code = {p.code.strip().upper(): p for p in partners if p.code}
        self.partners_by_name: dict[str, list[BusinessPartner]] = defaultdict(list)
        for p in partners:
            self.partners_by_name[p.name.strip().upper()].append(p)
        items = db.execute(select(InventoryItem.id, InventoryItem.item_code).where(
            InventoryItem.company_id == company_id)).all()
        self.items: dict[str, list[uuid.UUID]] = defaultdict(list)
        for item_id, code in items:
            self.items[(code or "").strip().upper()].append(item_id)
        self.locations = {(l.location or "").strip().upper(): l.id for l in db.scalars(
            select(ItemLocation).where(ItemLocation.company_id == company_id))}
        self.branches = {b.code.strip().upper(): b.id for b in db.scalars(
            select(Branch).where(Branch.company_id == company_id))}
        self.default_control = {}
        for control in SCHEDULE_CONTROL.values():
            candidates = [a for a in self.accounts.values()
                          if a.control_type == control and a.is_postable and a.is_active]
            self.default_control[control] = candidates[0] if len(candidates) == 1 else None

    def partner(self, code: str | None, name: str | None, role: str) -> tuple[BusinessPartner | None, str | None]:
        flag = "is_customer" if role == "customer" else "is_supplier"
        if code:
            p = self.partner_by_code.get(code.strip().upper())
            if p is None:
                return None, f"No {role} with code {code}."
            if not getattr(p, flag):
                return None, f"{p.code} {p.name} is not marked as a {role}."
            return p, None
        if name:
            matches = [p for p in self.partners_by_name.get(name.strip().upper(), [])
                       if getattr(p, flag)]
            if len(matches) == 1:
                return matches[0], None
            if len(matches) > 1:
                return None, f"{len(matches)} {role}s are called {name}; give the code."
            return None, f"No {role} called {name}."
        return None, f"The {role} is missing."


def resolve(db: Session, ctx: FinanceContext, batch: OpeningBalanceBatch) -> None:
    """Match every line's codes to master data and record per-row problems.
    Never creates master data. A missing customer is an error, not an insert."""
    look = _Lookups(db, ctx.company_id)
    seen: dict[tuple, int] = {}
    for line in batch.lines:
        errors = _parse_errors(line)
        schedule = line.schedule

        # branch
        line.branch_id = None
        if line.branch_code:
            line.branch_id = look.branches.get(line.branch_code.strip().upper())
            if line.branch_id is None:
                errors.append(f"No branch with code {line.branch_code}.")
        elif batch.branch_id:
            line.branch_id = batch.branch_id

        # account
        line.account_id = None
        control = SCHEDULE_CONTROL.get(schedule)
        if line.account_code:
            account = look.accounts.get(line.account_code.strip())
            if account is None:
                errors.append(f"No account with code {line.account_code}.")
            else:
                line.account_id = account.id
                if not account.is_postable:
                    errors.append(f"{account.code} {account.name} is a header account.")
                if not account.is_active:
                    errors.append(f"{account.code} {account.name} is inactive.")
                if control and account.control_type != control and _schedule_of(account.control_type) != schedule:
                    errors.append(f"{account.code} is not a {control.lower()} control account.")
                if account.control_type == "CUSTOMER_ADVANCE" and (line.amount or ZERO) > 0:
                    errors.append(f"{account.code} holds customer advances: enter the advance as a "
                                  "negative outstanding amount.")
        elif control:
            default = look.default_control.get(control)
            if default is None:
                errors.append(f"Give the account code: there is not exactly one "
                              f"{control.lower()} control account to default to.")
            else:
                line.account_id = default.id
                line.account_code = default.code
        else:
            errors.append("The account code is missing.")

        if schedule == "GL":
            debit, credit = line.debit or ZERO, line.credit or ZERO
            if debit < 0 or credit < 0:
                errors.append("Debit and credit cannot be negative; put the amount in the other column.")
            if debit > 0 and credit > 0:
                errors.append("Enter a debit or a credit, not both.")
            if debit == 0 and credit == 0:
                errors.append("The row has no amount.")
            key = ("GL", line.account_id, line.branch_id)

        elif schedule in ("AR", "AP"):
            role = "customer" if schedule == "AR" else "supplier"
            partner, problem = look.partner(line.party_code, line.party_name, role)
            line.customer_id = partner.id if partner and schedule == "AR" else None
            line.supplier_id = partner.id if partner and schedule == "AP" else None
            if problem:
                errors.append(problem)
            if not line.reference:
                errors.append("The invoice / bill reference is missing.")
            if line.amount is None or line.amount == 0:
                errors.append("The outstanding amount is missing or zero.")
            if line.document_date and line.document_date > batch.cutover_date:
                errors.append("The document is dated after the cut-over date.")
            if line.document_date and line.due_date and line.due_date < line.document_date:
                errors.append("The due date is before the document date.")
            key = (schedule, partner.id if partner else line.party_code, (line.reference or "").upper())

        elif schedule == "INVENTORY":
            ids = look.items.get((line.item_code or "").strip().upper(), [])
            line.inventory_item_id = ids[0] if len(ids) == 1 else None
            if not line.item_code:
                errors.append("The item code is missing.")
            elif not ids:
                errors.append(f"No inventory item with code {line.item_code}.")
            elif len(ids) > 1:
                errors.append(f"{len(ids)} items share code {line.item_code}.")
            line.location_id = None
            if line.location_name:
                line.location_id = look.locations.get(line.location_name.strip().upper())
                if line.location_id is None:
                    errors.append(f"No warehouse / location called {line.location_name}.")
            qty, cost = line.quantity, line.unit_cost
            if qty is None or qty <= 0:
                errors.append("The quantity must be greater than zero.")
            if cost is None or cost <= 0:
                errors.append("The unit cost is missing or zero; the valuation must give a cost.")
            if qty and cost and qty > 0 and cost > 0:
                computed = money(qty * cost)
                if line.amount is None:
                    line.amount = computed
                elif abs(line.amount - computed) > Decimal("0.01"):
                    errors.append(f"Total value {line.amount:,.2f} does not equal quantity × unit cost "
                                  f"({computed:,.2f}).")
            key = ("INVENTORY", line.inventory_item_id or line.item_code, line.location_id)

        else:  # BANK
            if not line.bank_name and not line.bank_account_number:
                errors.append("Give the bank and account number.")
            if line.statement_balance is None:
                errors.append("The statement balance is missing.")
            line.amount = ((line.statement_balance or ZERO) + (line.uncleared_deposits or ZERO)
                           - (line.uncleared_payments or ZERO))
            key = ("BANK", line.account_id)

        if key in seen:
            errors.append(f"Duplicate of row {seen[key]}.")
        else:
            seen[key] = line.row_number
        line.errors = errors
    with db_guard(db):
        db.flush()


def _parse_errors(line: OpeningBalanceLine) -> list[str]:
    """Errors from reading the file (kept across re-resolution). They are
    the ones written as 'Header: problem'."""
    raw_headers = {str(k) for k in (line.raw or {})}
    return [e for e in (line.errors or []) if e.split(":", 1)[0] in raw_headers]


def signed_amount(line: OpeningBalanceLine) -> Decimal:
    """Debit-positive value of a line for its account."""
    if line.schedule == "GL":
        return (line.debit or ZERO) - (line.credit or ZERO)
    if line.schedule == "AP":
        return -(line.amount or ZERO)
    return line.amount or ZERO


def _drcr(value: Decimal) -> str:
    """100,000.00 Cr rather than -100,000.00, as an accountant reads it."""
    return f"{abs(value):,.2f} {'Cr' if value < 0 else 'Dr'}"


def tie_out(db: Session, batch: OpeningBalanceBatch) -> list[dict]:
    """Trial-balance figure against schedule total, per control account."""
    accounts = {a.id: a for a in db.scalars(select(Account).where(Account.company_id == batch.company_id))}
    gl, sub = defaultdict(lambda: ZERO), defaultdict(lambda: ZERO)
    sub_schedule = {}
    for line in batch.lines:
        if not line.account_id:
            continue
        account = accounts.get(line.account_id)
        if line.schedule == "GL":
            if account and account.control_type in ALL_CONTROLS:
                gl[line.account_id] += signed_amount(line)
        else:
            sub[line.account_id] += signed_amount(line)
            sub_schedule[line.account_id] = line.schedule
    rows = []
    for account_id in sorted(set(gl) | set(sub), key=lambda a: accounts[a].code if a in accounts else ""):
        account = accounts.get(account_id)
        schedule = sub_schedule.get(account_id) or _schedule_of(account.control_type if account else None)
        rows.append({
            "account_id": str(account_id), "account_code": account.code if account else None,
            "account_name": account.name if account else None, "schedule": schedule,
            "trial_balance": gl[account_id], "schedule_total": sub[account_id],
            "difference": gl[account_id] - sub[account_id],
        })
    return rows


def validate(db: Session, ctx: FinanceContext, batch_id: uuid.UUID) -> dict:
    """Run every check. Clean → VALIDATED; otherwise stays DRAFT with the
    problems listed. The batch-level problems are returned; row problems are
    on the lines."""
    batch = get_batch(db, ctx, batch_id)
    _editable(batch)
    resolve(db, ctx, batch)
    problems: list[str] = []

    fy = db.get(FinancialYear, batch.financial_year_id)
    start, end = periods.fy_dates(fy)
    if not start <= batch.cutover_date <= end:
        problems.append(f"The cut-over date is outside {fy.name}.")
    try:
        period = periods.period_for_date(db, ctx.company_id, batch.cutover_date)
        if period.status != "OPEN":
            problems.append(f"Period {period.name} is {period.status}; open it before posting.")
    except FinanceError as error:
        problems.append(error.message)

    live = db.scalar(select(OpeningBalanceBatch.batch_number).where(
        OpeningBalanceBatch.company_id == ctx.company_id,
        OpeningBalanceBatch.financial_year_id == batch.financial_year_id,
        OpeningBalanceBatch.branch_id.is_(None) if batch.branch_id is None
        else OpeningBalanceBatch.branch_id == batch.branch_id,
        OpeningBalanceBatch.status.in_([OB_POSTED, OB_COMPLETED]),
        OpeningBalanceBatch.id != batch.id))
    if live:
        problems.append(f"Batch {live} is already posted for this year and branch.")

    gl_lines = [l for l in batch.lines if l.schedule == "GL"]
    if len(gl_lines) < 2:
        problems.append("The trial balance schedule needs at least two accounts.")
    total_dr = sum((l.debit or ZERO for l in gl_lines), ZERO)
    total_cr = sum((l.credit or ZERO for l in gl_lines), ZERO)
    batch.total_debit, batch.total_credit = total_dr, total_cr
    if total_dr != total_cr:
        problems.append(f"The trial balance does not balance: debits {total_dr:,.2f}, credits "
                        f"{total_cr:,.2f}, difference {abs(total_dr - total_cr):,.2f}.")

    for row in tie_out(db, batch):
        if row["difference"] != 0:
            label = f"{row['account_code']} {row['account_name']}"
            if row["schedule_total"] == 0:
                problems.append(f"{label} is {_drcr(row['trial_balance'])} in the trial balance but has "
                                f"no {row['schedule'] or 'subledger'} schedule behind it.")
            elif row["trial_balance"] == 0:
                problems.append(f"The {row['schedule']} schedule puts {_drcr(row['schedule_total'])} on "
                                f"{label}, which is not in the trial balance.")
            else:
                problems.append(f"{label}: trial balance {_drcr(row['trial_balance'])} but "
                                f"{row['schedule']} schedule {_drcr(row['schedule_total'])} "
                                f"(difference {abs(row['difference']):,.2f}).")

    row_errors = sum(1 for l in batch.lines if l.errors)
    if row_errors:
        problems.append(f"{row_errors} row(s) have problems; see the schedules.")

    passed = not problems
    batch.status = OB_VALIDATED if passed else OB_DRAFT
    batch.validated_at = datetime.now(timezone.utc) if passed else None
    batch.updated_by = ctx.user_id
    with db_guard(db):
        db.flush()
    record(db, ctx.audit, action="validate", module="finance", entity="opening_balance_batches",
           entity_id=batch.id, company_id=ctx.company_id,
           new_values={"passed": passed, "problems": len(problems), "row_errors": row_errors})
    return {"passed": passed, "problems": problems, "row_errors": row_errors}


# ------------------------------------------------------------ journal

def compose_lines(db: Session, batch: OpeningBalanceBatch) -> list[PostingLine]:
    """The opening journal: plain trial-balance lines, with each control
    account replaced by its schedule detail."""
    accounts = {a.id: a for a in db.scalars(select(Account).where(Account.company_id == batch.company_id))}
    covered = ALL_CONTROLS
    out: list[PostingLine] = []

    def add(account_id, value: Decimal, description: str, **extra):
        if value == 0:
            return
        out.append(PostingLine(
            account_id=account_id,
            debit=value if value > 0 else ZERO, credit=-value if value < 0 else ZERO,
            description=description, **extra))

    for line in batch.lines:
        account = accounts.get(line.account_id)
        if account is None:
            continue
        value = signed_amount(line)
        if line.schedule == "GL":
            if account.control_type in covered:
                continue
            add(account.id, value, f"Opening balance{': ' + line.reference if line.reference else ''}",
                branch_id=line.branch_id)
        elif line.schedule == "AR":
            add(account.id, value, f"Opening receivable {line.reference}",
                customer_id=line.customer_id, branch_id=line.branch_id)
        elif line.schedule == "AP":
            add(account.id, value, f"Opening payable {line.reference}",
                supplier_id=line.supplier_id, branch_id=line.branch_id)
        elif line.schedule == "INVENTORY":
            add(account.id, value, f"Opening stock {line.item_code}",
                inventory_item_id=line.inventory_item_id, quantity=line.quantity,
                branch_id=line.branch_id)
        elif line.schedule == "BANK":
            add(account.id, value,
                f"Opening bank balance {line.bank_name or ''} {line.bank_account_number or ''}".strip(),
                branch_id=line.branch_id)
    return out


def preview(db: Session, ctx: FinanceContext, batch_id: uuid.UUID) -> dict:
    batch = get_batch(db, ctx, batch_id)
    lines = compose_lines(db, batch)
    accounts = {a.id: a for a in db.scalars(select(Account).where(Account.company_id == ctx.company_id))}
    partners = {p.id: p.name for p in db.scalars(select(BusinessPartner).where(
        BusinessPartner.id.in_({x for l in lines for x in (l.customer_id, l.supplier_id) if x})))}
    total_dr = sum((l.debit for l in lines), ZERO)
    total_cr = sum((l.credit for l in lines), ZERO)
    schedules = {}
    for schedule in SCHEDULES:
        rows = [l for l in batch.lines if l.schedule == schedule]
        schedules[schedule] = {"rows": len(rows), "errors": sum(1 for l in rows if l.errors),
                               "total": sum((signed_amount(l) for l in rows), ZERO)}
    return {
        "batch_id": str(batch.id), "status": batch.status,
        "lines": [{
            "account_code": accounts[l.account_id].code, "account_name": accounts[l.account_id].name,
            "description": l.description, "debit": l.debit, "credit": l.credit,
            "party": partners.get(l.customer_id or l.supplier_id),
            "quantity": l.quantity,
        } for l in lines],
        "total_debit": total_dr, "total_credit": total_cr,
        "balanced": total_dr == total_cr and total_dr > 0,
        "schedules": schedules,
        "tie_out": tie_out(db, batch),
    }


def approve(db: Session, ctx: FinanceContext, batch_id: uuid.UUID) -> OpeningBalanceBatch:
    batch = get_batch(db, ctx, batch_id)
    if batch.status != OB_VALIDATED:
        raise FinanceError(f"Batch {batch.batch_number} is {batch.status}; validate it before approval.")
    # Master data may have changed since validation, so check again.
    result = validate(db, ctx, batch.id)
    if not result["passed"]:
        raise FinanceError("The batch no longer passes validation: " + "; ".join(result["problems"][:3]))
    batch.status = OB_APPROVED
    batch.approved_by = ctx.user_id
    batch.approved_at = datetime.now(timezone.utc)
    batch.updated_by = ctx.user_id
    with db_guard(db):
        db.flush()
    record(db, ctx.audit, action="approve", module="finance", entity="opening_balance_batches",
           entity_id=batch.id, company_id=ctx.company_id,
           old_values={"status": OB_VALIDATED}, new_values={"status": OB_APPROVED})
    return batch


def post(db: Session, ctx: FinanceContext, batch_id: uuid.UUID) -> OpeningBalanceBatch:
    batch = get_batch(db, ctx, batch_id)
    if batch.status != OB_APPROVED:
        raise FinanceError(f"Batch {batch.batch_number} is {batch.status}; it must be approved first.")
    lines = compose_lines(db, batch)
    journal = post_event(db, PostingRequest(
        event_code="OPENING_BALANCE_POSTED", company_id=ctx.company_id,
        journal_date=batch.cutover_date, branch_id=batch.branch_id,
        description=f"Opening balances at {batch.cutover_date:%d-%b-%Y} ({batch.batch_number})",
        reference=batch.source_document, document_type="OPENING_BALANCE_BATCH",
        document_id=batch.id, source_module="finance", source_reference=batch.batch_number,
        lines=lines, journal_type=OPENING,
    ), ctx.audit)
    # Customer lines become receivable items, so receipts can be allocated
    # against opening invoices exactly like new ones.
    from app.receivables.subledger import record_opening_items
    record_opening_items(db, batch, journal, ctx.audit)

    batch.status = OB_POSTED
    batch.journal_entry_id = journal.id
    batch.posted_by = ctx.user_id
    batch.posted_at = datetime.now(timezone.utc)
    batch.updated_by = ctx.user_id
    with db_guard(db):
        db.flush()
    record(db, ctx.audit, action="post", module="finance", entity="opening_balance_batches",
           entity_id=batch.id, company_id=ctx.company_id,
           old_values={"status": OB_APPROVED},
           new_values={"status": OB_POSTED, "journal_number": journal.journal_number,
                       "total_debit": journal.total_debit})
    return batch


def cancel(db: Session, ctx: FinanceContext, batch_id: uuid.UUID, reason: str) -> OpeningBalanceBatch:
    batch = get_batch(db, ctx, batch_id)
    if batch.status not in (OB_DRAFT, OB_VALIDATED, OB_APPROVED):
        raise FinanceError(f"Batch {batch.batch_number} is {batch.status} and cannot be cancelled.")
    if not reason or len(reason.strip()) < 3:
        raise FinanceError("Give a reason for cancelling the batch.")
    previous = batch.status
    batch.status = OB_CANCELLED
    batch.cancelled_by = ctx.user_id
    batch.cancelled_at = datetime.now(timezone.utc)
    batch.updated_by = ctx.user_id
    with db_guard(db):
        db.flush()
    record(db, ctx.audit, action="cancel", module="finance", entity="opening_balance_batches",
           entity_id=batch.id, company_id=ctx.company_id,
           old_values={"status": previous}, new_values={"status": OB_CANCELLED}, reason=reason.strip())
    return batch


# ------------------------------------------------------------ reconciliation

CATEGORY_LABELS = {
    "TRIAL_BALANCE": "Trial balance (total debits)",
    "ACCOUNTS_RECEIVABLE": "Accounts receivable",
    "ACCOUNTS_PAYABLE": "Accounts payable",
    "INVENTORY": "Inventory",
    "CASH": "Cash",
    "BANK": "Bank",
    "EQUITY": "Equity",
}


def erp_balances(db: Session, ctx: FinanceContext, batch: OpeningBalanceBatch) -> dict[str, Decimal]:
    """What the posted ledger says at the cut-over date, per category, each
    on its natural side (liabilities and equity as positive credits)."""
    tb = ledger.trial_balance(db, ctx, date_to=batch.cutover_date, branch_id=batch.branch_id)
    control = defaultdict(lambda: ZERO)
    equity = ZERO
    accounts = {a.id: a for a in db.scalars(select(Account).where(Account.company_id == ctx.company_id))}
    for row in tb["rows"]:
        account = accounts[uuid.UUID(row["account_id"])]
        if account.control_type:
            control[account.control_type] += row["balance"]
        if row["account_type"] == "EQUITY":
            equity += row["balance"]
    return {
        "TRIAL_BALANCE": tb["totals"]["closing_debit"],
        "ACCOUNTS_RECEIVABLE": control["RECEIVABLE"],
        "ACCOUNTS_PAYABLE": -control["PAYABLE"],
        "INVENTORY": control["INVENTORY"],
        "CASH": control["CASH"],
        "BANK": control["BANK"],
        "EQUITY": -equity,
    }


def subledger_checks(db: Session, ctx: FinanceContext, batch: OpeningBalanceBatch) -> list[dict]:
    """Receivable and payable control balances against the sum of their
    customer / supplier lines. They agree by construction; this proves it."""
    rows = []
    for control, party_col, label in (("RECEIVABLE", JournalEntryLine.customer_id, "Customers"),
                                      ("PAYABLE", JournalEntryLine.supplier_id, "Suppliers")):
        base = (select(func.coalesce(func.sum(JournalEntryLine.base_debit - JournalEntryLine.base_credit), 0))
                .join(JournalEntry, JournalEntry.id == JournalEntryLine.journal_entry_id)
                .join(Account, Account.id == JournalEntryLine.account_id)
                .where(JournalEntry.company_id == ctx.company_id,
                       JournalEntry.status.in_(ledger.LEDGER_STATUSES),
                       JournalEntry.journal_date <= batch.cutover_date,
                       Account.control_type == control))
        control_total = Decimal(db.scalar(base) or 0)
        party_total = Decimal(db.scalar(base.where(party_col.is_not(None))) or 0)
        rows.append({"subledger": label, "control_account_balance": control_total,
                     "subledger_total": party_total, "difference": control_total - party_total})
    return rows


def reconciliation(db: Session, ctx: FinanceContext, batch_id: uuid.UUID) -> dict:
    batch = get_batch(db, ctx, batch_id)
    saved = {r.category: r for r in db.scalars(select(OpeningBalanceReconciliation).where(
        OpeningBalanceReconciliation.batch_id == batch.id))}
    posted = batch.status in (OB_POSTED, OB_COMPLETED)
    erp = erp_balances(db, ctx, batch) if posted else {}
    rows = []
    for category in RECONCILIATION_CATEGORIES:
        entry = saved.get(category)
        erp_value = erp.get(category, ZERO)
        legacy = entry.legacy_balance if entry else None
        difference = (erp_value - legacy) if legacy is not None else None
        rows.append({
            "category": category, "label": CATEGORY_LABELS[category],
            "legacy_balance": legacy, "erp_balance": erp_value, "difference": difference,
            "explanation": entry.explanation if entry else None,
            "status": ("missing" if legacy is None else "agreed" if difference == 0
                       else "explained" if (entry.explanation or "").strip() else "unexplained"),
        })
    blocking = [r for r in rows if r["status"] in ("missing", "unexplained")]
    return {
        "batch_id": str(batch.id), "status": batch.status, "cutover_date": batch.cutover_date,
        "available": posted, "rows": rows,
        "subledgers": subledger_checks(db, ctx, batch) if posted else [],
        "can_complete": batch.status == OB_POSTED and not blocking,
        "blocking": [r["label"] for r in blocking],
    }


def save_reconciliation(db: Session, ctx: FinanceContext, batch_id: uuid.UUID,
                        entries: list[dict]) -> dict:
    batch = get_batch(db, ctx, batch_id)
    if batch.status != OB_POSTED:
        raise FinanceError("Reconciliation is entered after posting and before completion.")
    erp = erp_balances(db, ctx, batch)
    saved = {r.category: r for r in db.scalars(select(OpeningBalanceReconciliation).where(
        OpeningBalanceReconciliation.batch_id == batch.id))}
    for data in entries:
        category = data["category"]
        if category not in RECONCILIATION_CATEGORIES:
            raise FinanceError(f"Unknown reconciliation category {category}.")
        entry = saved.get(category)
        if entry is None:
            entry = OpeningBalanceReconciliation(batch_id=batch.id, category=category)
            db.add(entry)
        before = {"legacy_balance": entry.legacy_balance, "explanation": entry.explanation}
        legacy = data.get("legacy_balance")
        entry.legacy_balance = money(legacy) if legacy not in (None, "") else None
        entry.erp_balance = erp[category]
        entry.difference = (erp[category] - entry.legacy_balance) if entry.legacy_balance is not None else None
        entry.explanation = (data.get("explanation") or "").strip() or None
        entry.updated_by = ctx.user_id
        record(db, ctx.audit, action="update", module="finance",
               entity="opening_balance_reconciliations", entity_id=batch.id,
               company_id=ctx.company_id, old_values=before,
               new_values={"category": category, "legacy_balance": entry.legacy_balance,
                           "erp_balance": entry.erp_balance, "explanation": entry.explanation})
    with db_guard(db):
        db.flush()
    return reconciliation(db, ctx, batch.id)


def complete(db: Session, ctx: FinanceContext, batch_id: uuid.UUID) -> OpeningBalanceBatch:
    """Close the opening process. Refused while any category lacks its
    legacy figure or carries an unexplained difference."""
    state = reconciliation(db, ctx, batch_id)
    batch = get_batch(db, ctx, batch_id)
    if batch.status != OB_POSTED:
        raise FinanceError(f"Batch {batch.batch_number} is {batch.status}; only a posted batch can be completed.")
    if not state["can_complete"]:
        raise FinanceError("The opening position cannot be completed while these are unreconciled: "
                           + ", ".join(state["blocking"]) + ".")
    batch.status = OB_COMPLETED
    batch.completed_by = ctx.user_id
    batch.completed_at = datetime.now(timezone.utc)
    batch.updated_by = ctx.user_id
    with db_guard(db):
        db.flush()
    record(db, ctx.audit, action="approve", module="finance", entity="opening_balance_batches",
           entity_id=batch.id, company_id=ctx.company_id,
           old_values={"status": OB_POSTED}, new_values={"status": OB_COMPLETED},
           reason="Opening balances reconciled and locked")
    return batch


# ------------------------------------------------------------ template

def template_workbook() -> bytes:
    """The workbook the accountant fills in: one sheet per schedule plus notes."""
    workbook = Workbook()
    notes = workbook.active
    notes.title = "Read Me"
    for row in [
        ["Jezseem ERP — opening balance template"],
        [],
        ["Trial Balance", "The approved legacy closing trial balance at the cut-over date. Must balance."],
        ["Customers", "Every unpaid customer invoice. Total must equal the receivable account(s) in the trial balance."],
        ["Suppliers", "Every unpaid supplier bill. Total must equal the payable account(s)."],
        ["Inventory", "Item-wise valuation: quantity × unit cost. Total must equal the inventory account."],
        ["Banks", "Per bank account: statement balance + uncleared deposits − uncleared payments "
                  "must equal the bank account in the trial balance."],
        [],
        ["Codes", "Account, customer, supplier, item and branch codes must already exist in the ERP. "
                  "The import never creates master data."],
        ["Amounts", "Plain numbers. A negative outstanding amount is an advance / credit balance."],
        ["Dates", "DD-MM-YYYY."],
    ]:
        notes.append(row)
    notes["A1"].font = Font(bold=True, size=13)
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 110

    for schedule, columns in TEMPLATE_COLUMNS.items():
        sheet = workbook.create_sheet(SHEETS[schedule])
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="FFEFF4EF")
        for index in range(1, len(columns) + 1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = 20
        sheet.freeze_panes = "A2"

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def batch_summary(batch: OpeningBalanceBatch) -> dict:
    counts = defaultdict(int)
    errors = defaultdict(int)
    for line in batch.lines:
        counts[line.schedule] += 1
        if line.errors:
            errors[line.schedule] += 1
    return {"lines": dict(counts), "errors": dict(errors)}
