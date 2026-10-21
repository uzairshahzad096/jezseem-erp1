"""Stock movement analysis: repeat sellers, dead stock and restock."""

import io
from datetime import date, timedelta
from decimal import Decimal as D

from openpyxl import load_workbook

from app.inventory import stock
from app.inventory.models import InventoryItem
from app.reports import movement
from app.reports.stock_workbook import build_workbook
from app.sales.models import QUOTATION, SALES_ORDER, SalesDocument, SalesDocumentLine

TODAY = date(2026, 10, 9)


def item(world, code, on_hand=0, price=100):
    i = InventoryItem(company_id=world.company.id, item_code=code, description=f"Item {code}", price=D(price))
    world.db.add(i)
    world.db.flush()
    if on_hand:
        stock.post_movement(world.db, company_id=world.company.id, item_id=i.id, direction=stock.IN,
                            quantity=D(on_hand), source_type="opening_balance", entry_date=date(2025, 1, 1))
    return i


def move(world, product, direction, qty, on, source_type="adjustment", source_id=None):
    stock.post_movement(world.db, company_id=world.company.id, item_id=product.id, direction=direction,
                        quantity=D(qty), source_type=source_type, source_id=source_id, entry_date=on)
    world.db.flush()


def sell(world, product, on, qty=1, status="Posted", doc_type=SALES_ORDER):
    doc = SalesDocument(company_id=world.company.id, doc_type=doc_type, doc_number=f"S{product.item_code}{on:%y%m%d}{qty}",
                        doc_date=on, customer_id=world.customer.id, status=status)
    doc.lines = [SalesDocumentLine(line_number=1, item_id=product.id, description="x", quantity=D(qty),
                                   unit_price=D(10), taxable_amount=D(10 * qty))]
    world.db.add(doc)
    world.db.flush()


def analyse(world, **kw):
    cfg = movement.Settings(period_from=date(2026, 7, 1), period_to=TODAY, period_label="FY", **kw)
    rows = {r["item_code"]: r for r in movement.analyse(world.db, world.company.id, cfg, today=TODAY)}
    return cfg, rows


def test_classification_and_restock(world, db):
    repeat = item(world, "REP", on_hand=10)
    move(world, repeat, stock.OUT, 7, date(2026, 8, 20))           # 3 of 10 left: 30%
    for n, day in enumerate([5, 20, 40]):
        sell(world, repeat, date(2026, 7, 1) + timedelta(days=day), qty=10)
    occasional = item(world, "OCC", on_hand=50)
    sell(world, occasional, date(2026, 8, 1))
    slow = item(world, "SLO")                                      # never stocked, but sells
    sell(world, slow, date(2026, 3, 1))                            # sold, but before this FY
    dead = item(world, "DED", on_hand=30, price=500)
    sell(world, dead, date(2024, 6, 1))                            # not sold for over a year
    never = item(world, "NEW", on_hand=5)                          # in stock, never sold
    inactive = item(world, "INA")                                  # nothing at all
    unposted = item(world, "UNP", on_hand=10)
    sell(world, unposted, date(2026, 8, 1), status="Un-Posted")    # does not count
    quoted = item(world, "QUO", on_hand=10)
    sell(world, quoted, date(2026, 8, 1), doc_type=QUOTATION)      # quotations never count

    cfg, rows = analyse(world)
    assert {k: rows[k]["movement"] for k in rows} == {
        "REP": "REPEAT", "OCC": "OCCASIONAL", "SLO": "SLOW", "DED": "DEAD", "NEW": "DEAD",
        "INA": "INACTIVE", "UNP": "DEAD", "QUO": "DEAD"}
    assert rows["REP"]["orders_in_period"] == 3 and rows["REP"]["qty_sold_in_period"] == 30

    # Restock: 30% of the 10 added is left. Refill to the added level (10
    # beats two months' cover of 2.5/month = 5): order 10 − 3 = 7.
    assert (rows["REP"]["added_level"], rows["REP"]["remaining_pct"]) == (10, 30)
    assert rows["REP"]["restock"] and rows["REP"]["restock_priority"] == "HIGH"
    assert rows["REP"]["suggested_restock_qty"] == 7 and rows["REP"]["restock_reason"] == "30% of 10 left"
    # Never stocked but selling: out of stock, order at least one.
    assert rows["SLO"]["restock"] and rows["SLO"]["restock_priority"] == "LOW"
    assert rows["SLO"]["added_level"] is None and rows["SLO"]["suggested_restock_qty"] == 1
    assert rows["SLO"]["remaining_pct"] == 0 and rows["INA"]["remaining_pct"] == 0   # nothing left
    assert rows["OCC"]["remaining_pct"] == 100
    assert not rows["OCC"]["restock"]                               # 50 of 50 left
    assert not rows["INA"]["restock"] and not rows["DED"]["restock"]

    totals = movement.summary(list(rows.values()))
    assert totals["counts"]["DEAD"] == 4 and totals["dead_quantity"] == 55
    assert totals["dead_list_value"] == 30 * 500 + 5 * 100 + 10 * 100 + 10 * 100
    assert totals["restock_total"] == 2 and totals["top_repeat"][0]["item_code"] == "REP"


def test_services_are_never_restocked(world, db):
    from app.inventory.models import MainProduct
    services = MainProduct(company_id=world.company.id, description="SERVICES")
    db.add(services)
    db.flush()
    p = item(world, "SRV")
    p.main_product_id = services.id
    for day in (3, 9, 15):
        sell(world, p, date(2026, 7, day))
    _, rows = analyse(world)
    assert rows["SRV"]["movement"] == "REPEAT" and not rows["SRV"]["restock"]


def test_restock_percent_of_stock_added(world, db):
    p = item(world, "PCT", on_hand=5)                              # opening 5 in Jan 2025
    move(world, p, stock.IN, 95, date(2026, 9, 1), "goods_receipt")  # had 5, received 95: level 100
    move(world, p, stock.OUT, 55, date(2026, 9, 10))
    _, rows = analyse(world)
    assert (rows["PCT"]["added_level"], rows["PCT"]["remaining_pct"], rows["PCT"]["restock"]) == (100, 45, False)

    move(world, p, stock.OUT, 5, date(2026, 9, 20))                 # 40 left: exactly 40%
    _, rows = analyse(world)
    assert rows["PCT"]["restock"] and rows["PCT"]["suggested_restock_qty"] == 60
    assert rows["PCT"]["restock_priority"] == "LOW"                 # flagged even though it never sold
    assert rows["PCT"]["added_on"] == date(2026, 9, 1)


def test_reversed_receipt_is_not_stock_added(world, db):
    import uuid
    p = item(world, "REV", on_hand=10)
    receipt = uuid.uuid4()
    move(world, p, stock.IN, 90, date(2026, 9, 1), "goods_receipt", receipt)
    move(world, p, stock.OUT, 90, date(2026, 9, 2), "goods_receipt_reversal", receipt)
    _, rows = analyse(world)
    # The level is still the opening 10, so 10 left is 100%, not 10%.
    assert (rows["REV"]["added_level"], rows["REV"]["remaining_pct"], rows["REV"]["restock"]) == (10, 100, False)


def test_thresholds_are_adjustable(world, db):
    p = item(world, "TWO", on_hand=10)
    move(world, p, stock.OUT, 6, date(2026, 8, 1))                 # 4 of 10 left
    sell(world, p, date(2026, 7, 10))
    sell(world, p, date(2026, 7, 20))
    _, rows = analyse(world, repeat_min=2)
    assert rows["TWO"]["movement"] == "REPEAT" and rows["TWO"]["restock"]
    _, rows = analyse(world, repeat_min=2, restock_percent=30)
    assert not rows["TWO"]["restock"]


def test_period_choices(world, db):
    start, end, label = movement.resolve_period(db, world.company.id, "fy", None, None, today=TODAY)
    assert (start, end) == (date(2026, 7, 1), TODAY) and "2026-27" in label
    assert movement.resolve_period(db, world.company.id, "calendar", None, None, today=TODAY)[0] == date(2026, 1, 1)
    assert movement.resolve_period(db, world.company.id, "last12", None, None, today=TODAY)[0] == TODAY - timedelta(days=365)


def test_workbook_has_all_tabs(world, db):
    p = item(world, "WB1")
    for day in (3, 9, 15):
        sell(world, p, date(2026, 7, day))
    cfg, rows = analyse(world)
    data = build_workbook("Test Co", cfg, list(rows.values()), movement.summary(list(rows.values())))
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Restock list", "Repeat sellers", "Dead stock", "All products"]
    restock = wb["Restock list"]
    assert restock["A5"].value == "WB1" and restock.freeze_panes == "C5" and restock.auto_filter.ref


def test_report_and_workbook_endpoints(client, world, db, monkeypatch):
    import app.reports.router as reports_router
    monkeypatch.setattr(reports_router, "active_company_id", lambda db: world.company.id)
    p = item(world, "API")
    for day in (3, 9, 15):
        sell(world, p, date(2026, 7, day))
    data = client.get("/api/v1/reports/product-stock?restock=true&period=last12").json()
    assert data["summary"]["restock_total"] >= 1 and data["items"][0]["item_code"] == "API"
    assert client.get("/api/v1/reports/product-stock?movement=REPEAT").json()["total"] >= 1
    workbook = client.get("/api/v1/reports/product-stock/workbook?restock_percent=50")
    assert workbook.status_code == 200 and workbook.content[:2] == b"PK"
